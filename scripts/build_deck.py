"""
Build docs/Raqib_submission_deck.pptx from the organiser's template.

Usage: python scripts/build_deck.py --base <template-with-3-additional-slides.pptx>
The base is the organiser template after duplicating the "Additional Slide" twice (8 slides total).
"""
import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "screenshots"
NAVY, TEAL, INK, MUTED = RGBColor(0x0B, 0x3D, 0x5C), RGBColor(0x0F, 0x76, 0x6E), RGBColor(0x0F, 0x17, 0x2A), RGBColor(0x47, 0x55, 0x69)
CARD, LINE, RED, WHITE = RGBColor(0xF1, 0xF5, 0xF9), RGBColor(0xCB, 0xD5, 0xE1), RGBColor(0xB4, 0x23, 0x18), RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"


def text(slide, x, y, w, h, runs, size=11, color=INK, bold=False, anchor=MSO_ANCHOR.TOP, name=None):
    """runs: str | list of paragraphs; a paragraph is str or list of (text, {bold,color,size})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(3)
        bullet = False
        if isinstance(para, str) and para.startswith("• "):
            bullet, para = True, para[2:]
        segs = para if isinstance(para, list) else [(para, {})]
        if bullet:
            _bullet(p)
        for t, o in segs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name, f.size = FONT, Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", color)
    return tb


def _bullet(p):
    from pptx.oxml.ns import qn
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.16)))
    pPr.set("indent", str(-Inches(0.13)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu)


def card(slide, x, y, w, h, fill=CARD, line=None, radius=True, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def title(slide, t, sub=None):
    text(slide, 0.4, 0.68, 9.2, 0.42, t, size=20, color=NAVY, bold=True, name="Title")
    if sub:
        text(slide, 0.4, 1.08, 9.2, 0.3, sub, size=11, color=MUTED, name="Subtitle")


def crop_shot(name, left=300, top=0, right=None, bottom=None):
    src = SHOTS / name
    tag = f"{left}_{top}"
    im = Image.open(src)
    box = (left, top, right or im.width, bottom or im.height)
    out = SHOTS / f"crop_{tag}_{name}"
    im.crop(box).save(out)
    return out


def picture(slide, path, x, y, w=None, h=None, border=True):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return pic


def remove_textboxes(slide, keep_pictures=True):
    for shp in list(slide.shapes):
        if shp.shape_type == 17:  # text box
            shp._element.getparent().remove(shp._element)


# =====================================================================================
def cover(s):
    fills = {"Team Name :": "", "Team Leader Name :": "", "Team Size :": "",
             "Problem Statement :": "Problem 1 — Risk, Fraud and Regulatory Intelligence Copilot  ·  Raqib"}
    for shp in s.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip() in fills:
            val = fills[shp.text_frame.text.strip()]
            if val:
                p = shp.text_frame.paragraphs[0]
                r = p.add_run()
                r.text = " " + val
                src = p.runs[0].font
                r.font.name = src.name or FONT
                r.font.size = src.size
                r.font.bold = False
                r.font.color.rgb = TEAL


def problem_brief(s):
    remove_textboxes(s)
    title(s, "1 · Problem brief: financial-crime and risk work is manual and siloed",
          "Banking & NBFC (GCC) · AML/CFT, sanctions, fraud, Basel III liquidity, IFRS 9 credit risk")
    # left: problem, persona, context
    blocks = [
        ("Business problem", "Analysts triage monitoring alerts across core-banking screens, policy PDFs and spreadsheets, "
                             "write every STR from scratch, and build LCR / NPL reporting in Excel. Decisions are slow, "
                             "inconsistent and hard to evidence for the regulator."),
        ("Target users", "FCC analyst (triage and investigate) · MLRO (decide and file STRs) · Treasury / Risk (LCR, NPL) · "
                         "Internal Audit and Model Risk (oversight)."),
        ("Industry context", "UAE AML/CFT framework (Federal Decree-Law 20/2018), goAML STR filing, tipping-off prohibition, "
                             "internal AED 55,000 cash threshold, Basel III LCR ≥ 100%, IFRS 9 staging."),
    ]
    y = 1.5
    for head, body in blocks:
        card(s, 0.4, y, 4.35, 1.17, name=f"Card {head}")
        text(s, 0.55, y + 0.08, 4.1, 0.26, head, size=11.5, color=TEAL, bold=True)
        text(s, 0.55, y + 0.34, 4.1, 0.82, body, size=10, color=INK)
        y += 1.25
    # right: today vs with Raqib
    rows = [("Job", "Today", "With Raqib"),
            ("Alert triage", "5+ systems; policy looked up by hand", "One evidence view + the governing clause"),
            ("Explain a risk", "Analyst writes SQL or raises a ticket", "Plain-English question → governed answer"),
            ("STR drafting", "Written from scratch each time", "Cited draft, amounts verified in code"),
            ("LCR / NPL", "Excel pack, days later", "Live LCR trigger + NPL by sector"),
            ("Audit trail", "Emails and spreadsheets", "Every question and tool action logged")]
    tbl = s.shapes.add_table(len(rows), 3, Inches(4.95), Inches(1.5), Inches(4.65), Inches(3.67)).table
    tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(1.05), Inches(1.7), Inches(1.9)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            cell.margin_left = cell.margin_right = Inches(0.06)
            para = cell.text_frame.paragraphs[0]
            f = para.runs[0].font
            f.name, f.size = FONT, Pt(9.5 if r else 10)
            f.bold = r == 0 or c == 0
            f.color.rgb = WHITE if r == 0 else (TEAL if c == 2 else INK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (WHITE if r % 2 else CARD)


def architecture(s):
    title(s, "2 · Architecture: one governed platform, end to end")
    img = ROOT / "docs" / "architecture.png"
    h = 3.92
    w = h * 16 / 9
    picture(s, img, (10 - w) / 2, 1.16, w=w, h=h)
    text(s, 0.4, 5.1, 9.2, 0.25,
         "Data: structured core banking, digital, treasury and credit tables + unstructured policy PDFs. "
         "CoCo skills: synthetic-banking-data → aml-investigation → regulatory-report-drafter → copilot-eval → alert-dispatch.",
         size=9, color=MUTED)


def impact(s):
    title(s, "3 · Impact: measured on the synthetic bank, built to scale",
          "Measured = from the repo's automated tests and eval. Business outcomes are estimates to validate in a pilot.")
    stats = [("100%", "recall on 42 planted\ntypology scenarios (8 rules)"),
             ("≤ 12", "false-positive alerts per\nrule on 2,150 customers"),
             ("10 / 10", "golden copilot questions\npassed (incl. guardrails)"),
             ("3 / 3", "hallucination, tipping-off and\nunauthorised filing blocked")]
    x = 0.4
    for big, label in stats:
        card(s, x, 1.5, 2.17, 1.3, name=f"Stat {big}")
        text(s, x + 0.12, 1.58, 1.95, 0.55, big, size=26, color=TEAL, bold=True)
        text(s, x + 0.12, 2.15, 1.95, 0.6, label, size=9.5, color=INK)
        x += 2.34
    text(s, 0.4, 3.0, 4.45, 0.3, "Expected outcomes (estimates)", size=12, color=NAVY, bold=True)
    text(s, 0.4, 3.32, 4.45, 1.9, [
        "• Alert triage: evidence + policy in one screen, so far less time per alert spent switching between systems",
        "• STR first draft in under 2 minutes instead of hours; analysts review rather than write",
        "• Consistent answers: one governed metric definition used by every team and surface",
        "• Audit-ready by default: every number traceable to transactions and a policy clause",
    ], size=10)
    text(s, 5.15, 3.0, 4.45, 0.3, "Scalability & beyond the demo", size=12, color=NAVY, bold=True)
    text(s, 5.15, 3.32, 4.45, 1.9, [
        "• Snowflake-native: dynamic tables and Cortex scale with warehouse size, and data never leaves the platform",
        "• New typology = one SQL view + rulebook entry + test; ontology extends via the semantic views",
        "• Reusable CoCo skills and tools port to any bank: swap table names, keep the guardrails",
        "• Next: trade-finance (TBML) rules, ML anomaly scores, goAML XML export, Arabic UI",
    ], size=10)


def coco_lifecycle(s):
    remove_textboxes(s)
    title(s, "Additional · CoCo across the whole lifecycle",
          "Every phase runs as a CoCo prompt against Snowflake; sessions are the evidence (coco/EVIDENCE.md)")
    phases = [("PLAN", "Profile the data, frame personas, critique the data model / ontology / workflow, check rule ↔ policy thresholds",
               "coco/01_plan.md"),
              ("BUILD", "Generate and load data, then build dynamic tables, 8 rules, PDF → Cortex Search, tools, semantic views, agent and app, fixing errors in the repo",
               "coco/02_build.md"),
              ("EXECUTE", "Inject live activity → new alert → case → STR → Slack/Jira via MCP; schedule tasks and CoCo runs",
               "coco/03_execute.md"),
              ("TEST", "Recall and false-positive back-test, ontology consistency, masking / RBAC, guardrail and edge cases, golden eval",
               "coco/04_test.md")]
    x = 0.4
    for name, body, f in phases:
        card(s, x, 1.5, 2.17, 1.95, name=f"Phase {name}")
        text(s, x + 0.12, 1.58, 1.95, 0.3, name, size=13, color=TEAL, bold=True)
        text(s, x + 0.12, 1.9, 1.95, 1.2, body, size=9.5)
        text(s, x + 0.12, 3.12, 1.95, 0.25, f, size=8.5, color=MUTED)
        x += 2.34
    card(s, 0.4, 3.6, 4.5, 1.62, fill=WHITE, line=LINE, name="Fixes card")
    text(s, 0.55, 3.66, 4.25, 0.28, "Real fixes CoCo made on the live account", size=11.5, color=NAVY, bold=True)
    text(s, 0.55, 3.95, 4.25, 1.25, [
        "• Dynamic table over views → wrapped in DYNAMIC_TABLE_REFRESH_BOUNDARY",
        "• Semantic view metric rules → rewrote ECL / coverage metrics",
        "• No streams on FULL-refresh tables → poll-based alert router",
        "• Masking needs Enterprise → secure view with role check",
    ], size=9.5)
    card(s, 5.1, 3.6, 4.5, 1.62, fill=WHITE, line=LINE, name="Ingenuity card")
    text(s, 5.25, 3.66, 4.25, 0.28, "Ingenuity in CoCo usage", size=11.5, color=NAVY, bold=True)
    text(s, 5.25, 3.95, 4.25, 1.25, [
        [("5 skills, ", {"bold": True}), ("3 subagents with HANDOFF + independent validation", {})],
        [("Plan → build: ", {"bold": True}), ("CoCo's own review added the PIPELINE_HEALTH check", {})],
        [("MCP + schedules: ", {"bold": True}), ("Slack / Jira dispatch, Snowflake tasks", {})],
        [("6 custom tools: ", {"bold": True}), ("stored procedures the agent calls to act", {})],
    ], size=9.5)


def guardrails(s):
    remove_textboxes(s)
    title(s, "Additional · Governed, explainable, and safe by construction")
    rows = [("Risk", "Control (enforced in code / Snowflake)"),
            ("Hallucinated numbers", "Every AED amount reconciled to evidence; failing drafts cannot be filed"),
            ("Fake citations", "Cited chunk IDs must be among the retrieved policy passages"),
            ("Tipping-off", "Language filter + agent instructions; flagged text removed"),
            ("Unauthorised filing", "Agent has no filing tool; APPROVE_REPORT needs RAQIB_MLRO"),
            ("PII exposure", "Secure view masks Emirates ID / phone / email / DOB unless the MLRO role is active"),
            ("Model outage", "Primary → fallback model → deterministic template"),
            ("Unknown entities", "Explicit \"not found\" (tested), never a guess"),
            ("Oversight", "AUDIT_LOG + COPILOT_AUDIT + EVAL_RESULTS in Snowflake")]
    tbl = s.shapes.add_table(len(rows), 2, Inches(0.4), Inches(1.2), Inches(5.0), Inches(3.95)).table
    tbl.columns[0].width, tbl.columns[1].width = Inches(1.45), Inches(3.55)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            cell.margin_left = cell.margin_right = Inches(0.06)
            f = cell.text_frame.paragraphs[0].runs[0].font
            f.name, f.size, f.bold = FONT, Pt(9.5), (r == 0 or c == 0)
            f.color.rgb = WHITE if r == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (WHITE if r % 2 else CARD)
    img = crop_shot("deck_str.png", left=360, top=150, right=1380, bottom=700)
    picture(s, img, 5.6, 1.2, w=4.0)
    text(s, 5.6, 3.45, 4.0, 0.9, [[("Validated STR draft: ", {"bold": True}),
                                   ("8/8 amounts reconciled, citations valid, tipping-off check clear. Only the MLRO can approve "
                                    "and file; a goAML reference is recorded and every step is logged.", {})]], size=9.5, color=MUTED)


def walkthrough(s):
    remove_textboxes(s)
    title(s, "Additional · Product walkthrough (Streamlit in Snowflake)")
    left = SHOTS / "live_command_center.png"  # captured from Streamlit in Snowflake (live Cortex)
    im = Image.open(left)
    h = 3.55
    w = h * im.width / im.height
    picture(s, left, 0.4, 1.2, w=w, h=h)
    text(s, 0.4, 4.8, w, 0.4, "Live in Snowflake: Command Center reading the deployed pipeline (Live · Snowflake Cortex)", size=9.5, color=MUTED)
    rx, rw = 0.4 + w + 0.25, 9.6 - (0.4 + w + 0.25)
    y = 1.2
    for fname, box, cap in [("deck_copilot.png", (380, 395, 1380, 740), "Copilot: governed answer with rules, score breakdown and policy citations"),
                            ("deck_moneyflow.png", (360, 285, 1380, 560), "Investigate: why the alert fired, against the KYC profile and evidence")]:
        img = crop_shot(fname, *box)
        im = Image.open(img)
        hh = rw * im.height / im.width
        picture(s, img, rx, y, w=rw, h=hh)
        text(s, rx, y + hh + 0.02, rw, 0.3, cap, size=9.5, color=MUTED)
        y += hh + 0.42


def links(s):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(4.55), Inches(9.0), Inches(0.75))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x0B, 0x1E, 0x33)
    box.line.fill.background()
    text(s, 0.7, 4.6, 8.6, 0.65, [
        [("GitHub  ", {"bold": True, "color": RGBColor(0x7D, 0xD3, 0xFC)}),
         ("github.com/ipavanreddy/raqib-snowflake-copilot", {"color": WHITE})],
        [("Live app  ", {"bold": True, "color": RGBColor(0x7D, 0xD3, 0xFC)}),
         ("Streamlit in Snowflake · RAQIB.APP.RAQIB_APP (Snowsight login)", {"color": WHITE})],
    ], size=11)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True)
    ap.add_argument("--out", default=str(ROOT / "docs" / "Raqib_submission_deck.pptx"))
    a = ap.parse_args()
    prs = Presentation(a.base)
    sl = list(prs.slides)
    assert len(sl) == 8, f"expected 8 slides in base, got {len(sl)}"
    cover(sl[0])
    problem_brief(sl[1])
    architecture(sl[2])
    impact(sl[3])
    coco_lifecycle(sl[4])
    guardrails(sl[5])
    walkthrough(sl[6])
    links(sl[7])
    prs.save(a.out)
    print("saved", a.out)


if __name__ == "__main__":
    main()
